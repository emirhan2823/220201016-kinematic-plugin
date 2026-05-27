"""Generate the CENG 428 submission report as DOCX with embedded screenshots."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SS_DIR = r"C:\Users\emirh\Documents\student-com-kinematic-plugin\screenshots"

doc = Document()

# -- page margins --
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style = doc.styles["Normal"]
style.font.name = "Cambria"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

# ═══════════════ Title ═══════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run("CENG 428 – Character Animation:\nModeling, Simulation and Human\nMotion Control")
run.bold = True
run.font.size = Pt(26)
run.font.name = "Cambria"
p.paragraph_format.space_after = Pt(14)

p = doc.add_paragraph()
run = p.add_run("220201016 – Emirhan Yıldız")
run.italic = True
run.font.size = Pt(14)
run.font.name = "Cambria"
p.paragraph_format.space_after = Pt(20)

# ═══════════════ Abstract ═══════════════
doc.add_paragraph(
    "This work presents a kinematic motion control plugin for the N8RO simulation "
    "environment. The character’s motion is driven entirely by direct joint-angle "
    "manipulation — no forces, torques, or rigid-body dynamics are involved. "
    "Ten major human joints (ankles, knees, hips, shoulders, elbows) are controlled "
    "through Euler rotation overrides, with additional trunk and head joints for "
    "posture. A simplified Center of Mass (CoM) estimation provides balance-aware "
    "corrections during the walk cycle. Three distinct motions — walking, pushing, "
    "and climbing — are registered as runtime animation evaluators on the Nathan "
    "human model. The result was verified visually through the N8RO GLB viewer."
)

# ═══════════════ Section 1 ═══════════════
h = doc.add_heading("Joint Angle Manipulation in the N8RO GLB Environment", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "The plugin integrates with N8RO through the IAnimationModel interface. "
    "When the simulation runs, the host calls the registered animation evaluator "
    "each tick. The evaluator computes joint angles from the current simulation "
    "time and returns sparse Euler-angle overrides for each driven joint. "
    "The GLB viewer reflects these overrides in real time on the character mesh."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(f"{SS_DIR}\\walking.png", width=Inches(5.8))
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "Figure 1. The human character model visualized in the N8RO GLB environment. "
    "Joint-angle overrides from the plugin produce a walking pose with natural arm swing."
)
run.italic = True
run.font.size = Pt(9)
p.paragraph_format.space_after = Pt(16)

# ═══════════════ Section 2 ═══════════════
h = doc.add_heading("Selected Human Joints for Kinematic Motion Control", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "Ten primary joints were selected from the GLB skeleton for kinematic control. "
    "These joints correspond to the five bilateral pairs that define the major "
    "articulations of human locomotion. The Joint Editor panel in the GLB viewer "
    "lists the configured Nathan joints along with their skeleton node indices."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(f"{SS_DIR}\\joints.png", width=Inches(5.8))
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "Figure 2. Ten predefined human joints were selected for kinematic pose "
    "manipulation in the N8RO GLB environment."
)
run.italic = True
run.font.size = Pt(9)
p.paragraph_format.space_after = Pt(16)

# ═══════════════ Section 3 ═══════════════
h = doc.add_heading("Runtime Joint Control and Pose Update Status", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "During simulation, the host updates 10 loaded skin-joint eulerRotation values "
    "directly in memory. The status panel confirms that in-memory gITF pose targets "
    "are active, and joint overrides are applied without requiring per-update GLB "
    "file read/write operations."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(f"{SS_DIR}\\status.png", width=Inches(4.6))
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "Figure 3. Runtime output showing direct Euler rotation updates applied to the "
    "loaded GLB skin-joint structure during pose manipulation."
)
run.italic = True
run.font.size = Pt(9)
p.paragraph_format.space_after = Pt(16)

# ═══════════════ Section 4 ═══════════════
h = doc.add_heading("Center of Mass Estimation and Balance Correction", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "A simplified anthropometric model estimates the Center of Mass (CoM) from "
    "the current joint configuration. Segment mass fractions are assigned to the "
    "trunk (46%), pelvis (14%), head (8%), thighs (10% each), and shins (6% each). "
    "The CoM position is compared against the support polygon center — derived from "
    "smooth foot-contact weights — and the resulting lateral and fore-aft errors "
    "drive small corrections to trunk roll, trunk pitch, and pelvis lateral shift."
)

doc.add_paragraph(
    "Foot contact transitions use Hermite smoothStep interpolation rather than "
    "hard boolean thresholds, ensuring that the balance corrections vary "
    "continuously across the gait cycle without introducing visible trembling."
)

# ═══════════════ Section 5 ═══════════════
h = doc.add_heading("Registered Animation Motions", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "The plugin registers three animation evaluators on the "
    "animationModelNathanHuman model type:"
)

table2 = doc.add_table(rows=4, cols=3)
table2.style = "Light Shading Accent 1"
t2h = ["Animation Code", "Motion Type", "Key Characteristics"]
t2d = [
    ("Student CoM Walk",  "Bipedal walk",  "Sinusoidal hip/knee cycles, CoM balance, counter-swing arms"),
    ("Student CoM Push",  "Two-hand push", "Forward trunk lean, extended shoulder reach, rhythmic knee bend"),
    ("Student CoM Climb", "Step/climb",    "High alternating knee lifts, compensatory arm reach"),
]
for i, txt in enumerate(t2h):
    c = table2.rows[0].cells[i]
    c.text = txt
    for p2 in c.paragraphs:
        for r2 in p2.runs:
            r2.bold = True
            r2.font.size = Pt(10)
for r, row in enumerate(t2d):
    for c, val in enumerate(row):
        cell = table2.rows[r + 1].cells[c]
        cell.text = val
        for p2 in cell.paragraphs:
            for r2 in p2.runs:
                r2.font.size = Pt(10)

doc.add_paragraph("")

# ═══════════════ Section 6 ═══════════════
h = doc.add_heading("Walk Animation Pipeline", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph("The walk evaluator follows a fixed pipeline each tick:")

steps = [
    "Compute gait phase from simulation time (stepFrequency = 1.25 Hz)",
    "Generate nominal joint angles from sinusoidal functions of the phase",
    "Infer foot-contact weights using smoothStep transitions",
    "Estimate Center of Mass from the current pose",
    "Compute lateral and fore-aft balance error against support polygon",
    "Apply conservative trunk and pelvis corrections",
    "Output sparse joint-angle overrides (Euler radians) to the host",
]
for s in steps:
    doc.add_paragraph(s, style="List Number")

# ═══════════════ Section 7 ═══════════════
h = doc.add_heading("Build and Deployment", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "The project builds with MSVC v143+ and CMake 3.20+ using the provided "
    "build script. The resulting DLL links against the N8RO SDK headers and "
    "the arkheon-astlib, arkheon-astsim libraries. No external physics libraries "
    "are required — all joint computations are self-contained."
)

p = doc.add_paragraph()
run = p.add_run("Build command:  ")
run.bold = True
run = p.add_run(".\\scripts\\build.ps1")
run.font.name = "Consolas"
run.font.size = Pt(10)

p = doc.add_paragraph()
run = p.add_run("Deploy path:  ")
run.bold = True
run = p.add_run("C:\\N8RO\\userPlugins\\sim\\student-com-kinematic-plugin.dll")
run.font.name = "Consolas"
run.font.size = Pt(10)

# ═══════════════ Section 8: Source Code ═══════════════
doc.add_page_break()
h = doc.add_heading("Appendix — Source Code", level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph(
    "The complete plugin source code is listed below. The header file defines the "
    "plugin class, and the implementation file contains all three animation evaluators "
    "along with the CoM estimation and balance correction logic."
)

# --- Header file ---
h2 = doc.add_heading("StudentComKinematicPlugin.h", level=2)
for run in h2.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

HEADER_CODE = r"""C:\Users\emirh\Documents\student-com-kinematic-plugin\include\StudentComKinematicPlugin.h"""
IMPL_CODE = r"""C:\Users\emirh\Documents\student-com-kinematic-plugin\src\StudentComKinematicPlugin.cpp"""

with open(HEADER_CODE, "r", encoding="utf-8") as f:
    header_src = f.read()

p = doc.add_paragraph()
run = p.add_run(header_src)
run.font.name = "Consolas"
run.font.size = Pt(7.5)
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(12)

# --- Implementation file ---
h2 = doc.add_heading("StudentComKinematicPlugin.cpp", level=2)
for run in h2.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

with open(IMPL_CODE, "r", encoding="utf-8") as f:
    impl_src = f.read()

p = doc.add_paragraph()
run = p.add_run(impl_src)
run.font.name = "Consolas"
run.font.size = Pt(7.5)
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(12)

# ═══════════════ Save ═══════════════
out = r"C:\Users\emirh\Documents\student-com-kinematic-plugin\CENG428_Report_220201016_Emirhan_Yildiz.docx"
doc.save(out)
print(f"Report saved: {out}")
